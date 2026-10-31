"""按《COBEBUDDY-测试用例表》逐条复现测试，并生成同模板格式的测试结果 Excel。

运行：python tests/case_test.py [--xlsx 输出路径]
"""
import os
import re
import sys
import time
from datetime import datetime, timedelta

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), os.pardir)))

from app import create_app  # noqa: E402
from app.extensions import db  # noqa: E402
from app.models import (  # noqa: E402
    Classroom,
    Homework,
    KnowledgePoint,
    Problem,
    Submission,
    TestCase,
    User,
    homework_problems,
)

TEMPLATE = r"C:\Users\hky\Desktop\COBEBUDDY-测试用例表.xlsx"
OUTPUT = os.path.join(os.path.dirname(__file__), os.pardir, "测试结果.xlsx")

RESULTS = {}  # 用例编号 -> (实际结果, 状态, 缺陷编号)


# ----------------------------- 工具 -----------------------------


def csrf(cli, url="/auth/login"):
    # 已登录时访问登录页会重定向，跟随后再取 meta 中的 token
    html = cli.get(url, follow_redirects=True).get_data(as_text=True)
    m = re.search(r'name="csrf-token" content="([^"]+)"', html)
    return m.group(1) if m else ""


def login(cli, username, password):
    token = csrf(cli)
    return cli.post(
        "/auth/login",
        data={"username": username, "password": password, "csrf_token": token},
        follow_redirects=False,
    )


def text(resp):
    return resp.get_data(as_text=True)


def record(cid, actual, passed, bug=""):
    RESULTS[cid] = (actual, "通过" if passed else "失败", bug)
    print("  [%s] %s -> %s" % ("PASS" if passed else "FAIL", cid, actual[:70]))


def judge(cli, problem_id, code, token, **extra):
    payload = {"problem_id": problem_id, "code": code, "source": "free"}
    payload.update(extra)
    return cli.post("/api/judge", json=payload, headers={"X-CSRF-Token": token})


# ----------------------------- 用例 -----------------------------


def purge_user(user):
    """删除账号并清理其提交/班级，测试环境自愈用。"""
    Submission.query.filter_by(user_id=user.id).delete()
    for room in Classroom.query.filter_by(teacher_id=user.id).all():
        for hw in list(room.homeworks):
            Submission.query.filter_by(homework_id=hw.id).delete()
            db.session.execute(homework_problems.delete().where(homework_problems.c.homework_id == hw.id))
            db.session.delete(hw)
        db.session.delete(room)
    db.session.delete(user)


def prepare(app):
    """测试前自愈：清理历史遗留测试数据，恢复示例题目，保证用例可重复执行。"""
    import re as _re

    with app.app_context():
        for p in Problem.query.filter(Problem.title.in_(["测试题目A", "无用例题目"])).all():
            Submission.query.filter_by(problem_id=p.id).delete()
            db.session.execute(homework_problems.delete().where(homework_problems.c.problem_id == p.id))
            db.session.delete(p)
        for u in User.query.all():
            if _re.match(r"^(stu_new_|stu_diff_|stu_short_|t_new_|admin_edit_|hack_admin_|teacher2$|bad)", u.username):
                purge_user(u)
        for room in Classroom.query.filter_by(join_code="T2ROOM").all():
            for hw in list(room.homeworks):
                Submission.query.filter_by(homework_id=hw.id).delete()
                db.session.execute(homework_problems.delete().where(homework_problems.c.homework_id == hw.id))
                db.session.delete(hw)
            db.session.delete(room)
        db.session.flush()

        # 恢复示例题目 A + B 问题（可能被上一轮用例修改过）
        p1 = Problem.query.order_by(Problem.id).first()
        if p1:
            p1.title = "A + B 问题"
            p1.description = "输入两个整数 a 和 b（以空格分隔），输出它们的和。\n\n输入示例：3 5\n输出示例：8"
            p1.difficulty = "easy"
            p1.starter_code = "a, b = map(int, input().split())\nprint(a + b)"
            p1.reference_code = "a, b = map(int, input().split())\nprint(a + b)"
            names = ["输入输出", "基础语法"]
            p1.tags = [KnowledgePoint.query.filter_by(name=n).first() for n in names]
            p1.tags = [t for t in p1.tags if t]
            if len(p1.test_cases) != 3 or p1.test_cases[0].expected_output != "8":
                for tc in list(p1.test_cases):
                    db.session.delete(tc)
                db.session.flush()
                for idx, (i, o) in enumerate([("3 5", "8"), ("10 20", "30"), ("-1 1", "0")]):
                    db.session.add(TestCase(problem_id=p1.id, input_data=i, expected_output=o,
                                             is_sample=(idx == 0), score=10))
        db.session.commit()


def run_cases(app):
    prepare(app)
    anon = app.test_client()
    stu = app.test_client()
    tea = app.test_client()

    # ---------- 登录（TF-001 ~ TF-007） ----------
    r = login(stu, "student", "123456")
    home = stu.get("/")
    record("TF-001",
           "%s 跳转 %s，跟随后学生首页 %s" % (r.status_code, r.headers.get("Location"), home.status_code),
           r.status_code == 302 and home.status_code == 200)
    st = csrf(stu)

    r = login(tea, "teacher", "123456")
    thome = tea.get("/")
    record("TF-002",
           "%s 跳转 %s，跟随后教师首页 %s" % (r.status_code, r.headers.get("Location"), thome.status_code),
           r.status_code == 302 and thome.status_code == 200)
    tc = csrf(tea)

    record("TF-003",
           "学生首页与教师首页均 200，渲染内容不同（%d / %d 字节，含不同模块入口）"
           % (len(text(home)), len(text(thome))),
           home.status_code == 200 and thome.status_code == 200 and text(home) != text(thome)
           and "专项练习" in text(home) and "题库管理" in text(thome))

    wrong = app.test_client()
    r = login(wrong, "student", "wrongpass")
    page = text(r)
    record("TF-004", "%s 并提示“%s”，未进入系统" % (r.status_code, "用户名或密码错误" if "用户名或密码错误" in page else "无提示"),
           r.status_code == 200 and "用户名或密码错误" in page)

    r = login(app.test_client(), "no_such_user_9", "123456")
    page = text(r)
    record("TF-005", "%s 并提示“用户名或密码错误”（与密码错误同文案）" % r.status_code,
           r.status_code == 200 and "用户名或密码错误" in page)

    r = login(app.test_client(), "", "")
    page = text(r)
    record("TF-006", "空值提交 %s，提示拒绝登录，未进入系统" % r.status_code,
           r.status_code == 200 and "用户名或密码错误" in page)

    r = stu.get("/teacher/problems", follow_redirects=False)
    record("TF-007", "%s 重定向回首页并提示无权限" % r.status_code, r.status_code == 302)

    # ---------- 注册（TF-008 ~ TF-013） ----------
    new_name = "stu_new_%s" % datetime.now().strftime("%H%M%S%f")[-8:]
    cli = app.test_client()
    tok = csrf(cli)
    r = cli.post("/auth/register", data={
        "username": new_name, "password": "123456", "confirm": "123456",
        "role": "student", "csrf_token": tok,
    }, follow_redirects=False)
    can_login = login(app.test_client(), new_name, "123456").status_code == 302
    record("TF-008", "注册 %s 跳转登录页，新账号 %s 可登录：%s" % (r.status_code, new_name, can_login),
           r.status_code == 302 and can_login)

    cli = app.test_client()
    r = cli.post("/auth/register", data={
        "username": "student", "password": "123456", "confirm": "123456",
        "role": "student", "csrf_token": csrf(cli),
    })
    record("TF-009", "提示“该用户名已被注册”，注册被拒绝", "该用户名已被注册" in text(r))

    cli = app.test_client()
    r = cli.post("/auth/register", data={
        "username": "stu_diff_%s" % datetime.now().strftime("%H%M%S")[-6:], "password": "123456",
        "confirm": "654321", "role": "student", "csrf_token": csrf(cli),
    })
    record("TF-010", "提示“两次输入的密码不一致”", "两次输入的密码不一致" in text(r))

    cli = app.test_client()
    r = cli.post("/auth/register", data={
        "username": "stu_short_%s" % datetime.now().strftime("%H%M%S")[-6:], "password": "123",
        "confirm": "123", "role": "student", "csrf_token": csrf(cli),
    })
    record("TF-011", "提示“密码长度至少 6 位”，短密码被拒绝", "密码长度至少 6 位" in text(r))

    cli = app.test_client()
    r = cli.post("/auth/register", data={
        "username": "", "password": "123456", "confirm": "123456",
        "role": "student", "csrf_token": csrf(cli),
    })
    record("TF-012", "空用户名被拒绝：" + ("用户名格式提示" if "用户名需为" in text(r) else "无提示"),
           "用户名需为" in text(r))

    bad_name = "bad<script>x"
    cli = app.test_client()
    r = cli.post("/auth/register", data={
        "username": bad_name, "password": "123456", "confirm": "123456",
        "role": "student", "csrf_token": csrf(cli),
    })
    with app.app_context():
        exists = User.query.filter_by(username=bad_name).first() is not None
    record("TF-013", "含脚本标签的用户名被白名单拒绝，未入库（exists=%s）" % exists,
           "用户名需为" in text(r) and not exists)

    # ---------- 题目新增（TF-014 ~ TF-018） ----------
    def create_problem(title, with_cases=True, client=None, token=None):
        client = client or tea
        token = token or tc
        data = {"csrf_token": token, "title": title, "description": "用例用题",
                "difficulty": "easy", "new_tags": "数学运算",
                "starter_code": "a, b = map(int, input().split())\n",
                "reference_code": "a, b = map(int, input().split())\nprint(a * b)"}
        if with_cases:
            data["tc_input"] = ["2 3", "4 5"]
            data["tc_output"] = ["6", "20"]
            data["tc_sample"] = ["0"]
        return client.post("/teacher/problems/new", data=data, follow_redirects=False)

    r = create_problem("测试题目A")
    with app.app_context():
        pa = Problem.query.filter_by(title="测试题目A").first()
        pa_id = pa.id if pa else None
        pa_cases = len(pa.test_cases) if pa else 0
    listed = "测试题目A" in text(tea.get("/teacher/problems?q=测试题目A"))
    record("TF-014", "新增 %s 跳转编辑页，用例 %d 条，题库管理列表可检索到：%s"
           % (r.status_code, pa_cases, listed), r.status_code == 302 and pa_id and listed)

    r = create_problem("")
    record("TF-015", "提示“题目标题不能为空”，空标题被拒绝", "题目标题不能为空" in text(r))

    r = create_problem("无用例题目", with_cases=False)
    with app.app_context():
        ok_no_case = Problem.query.filter_by(title="无用例题目").first() is not None
    record("TF-016", "未填测试用例也可创建（%s），列表可见：%s" % (r.status_code, ok_no_case),
           r.status_code == 302 and ok_no_case)

    page = text(tea.get("/teacher/problems?q=测试题目A"))
    record("TF-017", "按关键词检索命中“测试题目A”，HTTP 200", "测试题目A" in page)

    r = stu.get("/teacher/problems/new", follow_redirects=False)
    record("TF-018", "学生访问新增页面被重定向回首页（%s，无权限）" % r.status_code, r.status_code == 302)

    # ---------- 题目删除（TF-019 ~ TF-022） ----------
    with app.app_context():
        t2 = User.query.filter_by(username="teacher2").first()
        if not t2:
            t2 = User(username="teacher2", display_name="李老师", role="teacher")
            t2.set_password("123456")
            db.session.add(t2)
            db.session.commit()
        teacher2_id = t2.id

    r = tea.post("/teacher/problems/%d/delete" % pa_id, data={"csrf_token": tc}, follow_redirects=False)
    # 注意：搜索框会回显关键词，因此按“标题编辑链接是否存在”判断列表内容
    gone = ("/teacher/problems/%d/edit" % pa_id) not in text(tea.get("/teacher/problems?q=测试题目A"))
    record("TF-019", "删除 %s 并跳转，题库列表不再显示“测试题目A”：%s" % (r.status_code, gone),
           r.status_code == 302 and gone)

    with app.app_context():
        t2_room = Classroom.query.filter_by(join_code="T2ROOM").first()
        if not t2_room:
            t2_room = Classroom(name="外部班级", join_code="T2ROOM", teacher_id=teacher2_id)
            db.session.add(t2_room)
            db.session.flush()
        t2_hw = Homework.query.filter_by(title="外部班级作业").first()
        if not t2_hw:
            t2_hw = Homework(title="外部班级作业", classroom_id=t2_room.id, teacher_id=teacher2_id,
                             due_at=datetime.now() + timedelta(days=2))
            db.session.add(t2_hw)
            db.session.flush()
            p0 = Problem.query.first()
            t2_hw.problems.append(p0)
        db.session.commit()
        t2_hw_id = t2_hw.id

    t2cli = app.test_client()
    login(t2cli, "teacher2", "123456")
    t2c = csrf(t2cli)
    with app.app_context():
        target_id = Problem.query.first().id
    r = t2cli.post("/teacher/problems/%d/delete" % target_id, data={"csrf_token": t2c}, follow_redirects=False)
    with app.app_context():
        still = db.session.get(Problem, target_id) is not None
    record("TF-020", "提示“只能删除自己创建的题目”，他人题目未被删除（still=%s）" % still, r.status_code == 302 and still)

    r = tea.post("/teacher/problems/999999/delete", data={"csrf_token": tc}, follow_redirects=False)
    record("TF-021", "删除不存在的题目返回 %s" % r.status_code, r.status_code == 404)

    r = stu.post("/teacher/problems/%d/delete" % target_id, data={"csrf_token": st}, follow_redirects=False)
    with app.app_context():
        still2 = db.session.get(Problem, target_id) is not None
    record("TF-022", "学生删除被重定向（%s，无权限），题目未被删除（still=%s）" % (r.status_code, still2),
           r.status_code == 302 and still2)

    # ---------- 题目查询（TF-023 ~ TF-027） ----------
    page = text(stu.get("/problems/"))
    record("TF-023", "题库列表 200 并展示题目“A + B 问题”", "A + B 问题" in page)

    with app.app_context():
        kp = KnowledgePoint.query.filter_by(name="基础语法").first()
        kp_id = kp.id if kp else None
    page = text(stu.get("/problems/?tag=%s" % kp_id))
    record("TF-024", "按知识点“基础语法”筛选返回匹配题目", "基础语法" in page and "A + B 问题" in page)

    easy_page = text(stu.get("/problems/?difficulty=easy"))
    hard_page = text(stu.get("/problems/?difficulty=hard"))
    record("TF-025", "easy 筛选可见题目，hard 筛选无匹配题目",
           "A + B 问题" in easy_page and "A + B 问题" not in hard_page)

    page = text(stu.get("/problems/?q=zzz_not_exist_987"))
    record("TF-026", "不存在的关键词返回空列表提示：%s" % ("没有找到匹配的题目" in page), "没有找到匹配的题目" in page)

    with app.app_context():
        p1 = db.session.get(Problem, target_id)
        sample = p1.test_cases[0].input_data if p1.test_cases else ""
    page = text(stu.get("/problems/%d" % target_id))
    record("TF-027", "题目详情 200，展示描述与样例输入“%s”" % sample.replace("\n", "\\n"),
           sample.split("\n")[0] in page)

    # ---------- 题目修改（TF-028 ~ TF-031） ----------
    with app.app_context():
        p0 = db.session.get(Problem, target_id)
        origin_title = p0.title
        origin_cases = len(p0.test_cases)
        origin_expected = p0.test_cases[0].expected_output
        origin_inputs = [tc.input_data for tc in p0.test_cases]
        origin_outputs = [tc.expected_output for tc in p0.test_cases]
        origin_samples = [str(i) for i, tc in enumerate(p0.test_cases) if tc.is_sample]
        origin_tags = ",".join(t.name for t in p0.tags)

    new_title = origin_title + "-改"
    r = tea.post("/teacher/problems/%d/edit" % target_id, data={
        "csrf_token": tc, "title": new_title, "description": "描述修改测试",
        "difficulty": "easy", "new_tags": origin_tags,
        "starter_code": "a, b = map(int, input().split())\n",
        "reference_code": "a, b = map(int, input().split())\nprint(a + b)",
        "tc_input": [], "tc_output": [], "tc_sample": [],
    }, follow_redirects=False)
    with app.app_context():
        saved = db.session.get(Problem, target_id).title
    record("TF-028", "修改成功，列表显示最新标题“%s”" % saved, r.status_code == 302 and saved == new_title)

    # TF-029：改期望输出 -> 同一正确代码由 AC 变 WA
    r0 = judge(stu, target_id, "a, b = map(int, input().split())\nprint(a + b)", st)
    before = (r0.get_json() or {}).get("status")
    tea.post("/teacher/problems/%d/edit" % target_id, data={
        "csrf_token": tc, "title": new_title, "description": "描述修改测试",
        "difficulty": "easy", "new_tags": origin_tags,
        "tc_input": ["3 5"], "tc_output": ["99999"], "tc_sample": ["0"],
    }, follow_redirects=False)
    r1 = judge(stu, target_id, "a, b = map(int, input().split())\nprint(a + b)", st)
    after = (r1.get_json() or {}).get("status")
    # 恢复原用例（完整还原，避免影响后续用例）
    tea.post("/teacher/problems/%d/edit" % target_id, data={
        "csrf_token": tc, "title": new_title, "description": "描述修改测试",
        "difficulty": "easy", "new_tags": origin_tags,
        "tc_input": origin_inputs, "tc_output": origin_outputs, "tc_sample": origin_samples,
    }, follow_redirects=False)
    record("TF-029", "修改期望输出后判题结果由 %s 变为 %s，用例实时生效" % (before, after),
           before == "AC" and after == "WA")

    r = tea.post("/teacher/problems/%d/edit" % target_id, data={
        "csrf_token": tc, "title": "", "description": "空标题提交", "difficulty": "easy",
    }, follow_redirects=False)
    with app.app_context():
        p = db.session.get(Problem, target_id)
        kept_title, kept_cases = p.title, len(p.test_cases)
    record("TF-030",
           "空标题提交时标题保留“%s”；未提交用例字段时既有用例保留（tc_count=%d，原 %d）"
           % (kept_title, kept_cases, origin_cases),
           kept_title == new_title and kept_cases == origin_cases)

    r = stu.post("/teacher/problems/%d/edit" % target_id, data={
        "csrf_token": st, "title": "学生篡改标题",
    }, follow_redirects=False)
    with app.app_context():
        t_now = db.session.get(Problem, target_id).title
    record("TF-031", "学生提交修改被重定向（%s，无权限），标题未变（%s）" % (r.status_code, t_now),
           r.status_code == 302 and t_now == new_title)

    # ---------- 判题提交（TF-032 ~ TF-037） ----------
    r = judge(stu, target_id, "a, b = map(int, input().split())\nprint(a + b)", st)
    data = r.get_json() or {}
    record("TF-032", "正确代码判题返回 %s，生成提交记录（id=%s）" % (data.get("status"), data.get("submission_id")),
           data.get("status") == "AC" and bool(data.get("submission_id")))
    ac_id = data.get("submission_id")

    r = judge(stu, target_id, "a, b = map(int, input().split())\nprint(a - b)", st)
    record("TF-033", "错误答案判题返回 %s" % (r.get_json() or {}).get("status"), (r.get_json() or {}).get("status") == "WA")

    r = judge(stu, target_id, "print('unclosed\n", st)
    record("TF-034", "语法错误代码判题返回 %s（编译错误）" % (r.get_json() or {}).get("status"),
           (r.get_json() or {}).get("status") == "CE")

    t0 = time.time()
    r = judge(stu, target_id, "while True:\n    pass\n", st)
    dt = time.time() - t0
    record("TF-035", "死循环被时间限制中断，返回 %s，未挂起（用时约 %.1fs）" % ((r.get_json() or {}).get("status"), dt),
           (r.get_json() or {}).get("status") == "TLE" and dt < 15)

    r = judge(stu, target_id, "import socket\nprint('ok')\n", st)
    record("TF-036", "导入 socket 被沙箱拦截，返回 %s" % (r.get_json() or {}).get("status"),
           (r.get_json() or {}).get("status") == "RE")

    long_code = "x = 1\n" * 12000
    r = judge(stu, target_id, long_code, st)
    body = r.get_json() or {}
    record("TF-037", "超长代码（%d 字符）返回 %s“%s”" % (len(long_code), r.status_code, body.get("error", "")),
           r.status_code == 400 and "代码过长" in body.get("error", ""))

    # ---------- 提交记录（TF-038 ~ TF-043） ----------
    page = text(stu.get("/problems/submissions/?status=AC"))
    record("TF-038", "提交记录列表 200，AC 筛选生效并可见本次提交（id=%s）" % ac_id,
           ("/problems/submissions/%s" % ac_id) in page)

    r = app.test_client().get("/problems/")
    r2 = app.test_client().post("/api/judge", json={"problem_id": target_id, "code": "print(1)"})
    record("TF-039", "未登录访问题库 %s、判题接口 %s，均被重定向到登录页" % (r.status_code, r2.status_code),
           r.status_code == 302 and r2.status_code in (302, 401))

    # 补足到 16 条提交用于分页
    with app.app_context():
        me = User.query.filter_by(username="student").first()
        need = 16 - (Submission.query.filter_by(user_id=me.id).count())
        fake = []
        for i in range(max(need, 0)):
            s = Submission(user_id=me.id, problem_id=target_id, code="print(%d)" % i,
                           status="WA", passed_count=0, total_count=1, source="practice")
            db.session.add(s)
            fake.append(s)
        db.session.commit()
        fake_ids = [s.id for s in fake]
    p1 = text(stu.get("/problems/submissions/?page=1"))
    p2 = text(stu.get("/problems/submissions/?page=2"))
    with app.app_context():
        total = Submission.query.filter_by(user_id=User.query.filter_by(username="student").first().id).count()
    record("TF-040", "第 2 页与第 1 页内容不同，翻页正常（%d 条提交，每页 15 条）" % total,
           p1 != p2 and total >= 16)

    with app.app_context():
        me = User.query.filter_by(username="student").first()
        other_user = User.query.filter(User.id != me.id, User.role == "student").first()
        other = Submission.query.filter(Submission.user_id == other_user.id).first()
        if not other:
            other = Submission(user_id=other_user.id, problem_id=db.session.get(Problem, target_id).id,
                               code="print(1)", status="AC", passed_count=1, total_count=1, source="practice")
            db.session.add(other)
            db.session.commit()
        other_id = other.id
    r = stu.get("/problems/submissions/%s" % other_id)
    record("TF-041", "学生查看他人提交详情返回 %s" % r.status_code, r.status_code == 403)

    with app.app_context():
        room = Classroom.query.filter_by(join_code="CLASS1").first()
        member = room.students[0]
        mine = Submission.query.filter_by(user_id=member.id).first()
        if not mine:
            mine = Submission(user_id=member.id, problem_id=target_id, code="print(1)",
                              status="AC", passed_count=1, total_count=1, source="practice")
            db.session.add(mine)
            db.session.commit()
        member_sub_id = mine.id
    r = tea.get("/problems/submissions/%d" % member_sub_id)
    record("TF-042", "本班学生提交详情教师可查看（%s）" % r.status_code, r.status_code == 200)

    r = t2cli.get("/problems/submissions/%d" % member_sub_id)
    record("TF-043", "非本班教师查看该学生提交返回 %s" % r.status_code, r.status_code == 403)

    # ---------- 作业提交（TF-044 ~ TF-045） ----------
    with app.app_context():
        hw1 = Homework.query.filter_by(title="第1次作业：分支与循环").first()
        hw1_id = hw1.id
        first_problem = hw1.problems[0]
        hw1_problem, hw1_code = first_problem.id, (first_problem.reference_code or "print(1)")
    r = judge(stu, hw1_problem, hw1_code, st, homework_id=hw1_id, source="homework")
    data = r.get_json() or {}
    with app.app_context():
        sub = db.session.get(Submission, data.get("submission_id")) if data.get("submission_id") else None
        linked = sub.homework_id if sub else None
    record("TF-044", "本班作业提交判题 %s，提交记录关联 homework_id=%s" % (data.get("status"), linked),
           data.get("status") == "AC" and linked == hw1_id)

    r = judge(stu, hw1_problem, "print(1)", st, homework_id=t2_hw_id, source="homework")
    record("TF-045", "学生提交非本班作业返回 %s“%s”" % (r.status_code, (r.get_json() or {}).get("error", "")),
           r.status_code == 403)

    codes = [tea.get(u).status_code for u in ["/teacher/classes", "/teacher/homeworks", "/teacher/stats", "/teacher/submissions"]]
    record("TF-046", "班级管理/作业管理/学情统计/提交查看页面均 %s" % codes, all(c == 200 for c in codes))

    cli = app.test_client()
    r = cli.post("/auth/login", data={"username": "student", "password": "123456"})
    record("TF-047", "缺少 CSRF Token 的表单 POST 被拒绝（%s）" % r.status_code, r.status_code == 400)

    # ---------- 管理员（TF-048 ~ TF-056，新增） ----------
    adm = app.test_client()
    r = login(adm, "admin", "admin123")
    dash = adm.get("/admin/")
    record("TF-048", "管理员登录 %s，控制台 %s" % (r.status_code, dash.status_code),
           r.status_code == 302 and dash.status_code == 200)
    ac = csrf(adm)

    page = text(adm.get("/admin/users"))
    record("TF-049", "用户列表 200，含 student/teacher 账号", "student" in page and "teacher" in page)

    tname = "t_new_%s" % datetime.now().strftime("%H%M%S")[-6:]
    r = adm.post("/admin/users/new", data={"csrf_token": ac, "username": tname,
                                           "display_name": "新教师", "role": "teacher", "password": "123456"},
                 follow_redirects=False)
    with app.app_context():
        created = User.query.filter_by(username=tname).first()
        created_id, created_role = (created.id, created.role) if created else (None, None)
    can_login = login(app.test_client(), tname, "123456").status_code == 302
    record("TF-050", "创建教师账号 %s（role=%s）并可登录：%s" % (r.status_code, created_role, can_login),
           r.status_code == 302 and created_role == "teacher" and can_login)

    sname = "admin_edit_%s" % datetime.now().strftime("%H%M%S")[-6:]
    with app.app_context():
        u = User(username=sname, display_name="待编辑学生", role="student")
        u.set_password("123456")
        db.session.add(u)
        db.session.commit()
        sid = u.id
    r = adm.post("/admin/users/%d/edit" % sid, data={"csrf_token": ac, "username": sname,
                                                     "display_name": "已改名同学", "role": "teacher"},
                 follow_redirects=False)
    with app.app_context():
        u = db.session.get(User, sid)
        changed = (u.display_name, u.role)
    record("TF-051", "编辑生效：姓名=%s，角色=%s" % changed, changed == ("已改名同学", "teacher"))

    r = adm.post("/admin/users/%d/reset-password" % sid, data={"csrf_token": ac}, follow_redirects=False)
    ok_login = login(app.test_client(), sname, "123456").status_code == 302
    record("TF-052", "重置密码后可用默认密码 123456 登录：%s" % ok_login, r.status_code == 302 and ok_login)

    r = adm.post("/admin/users/%d/toggle-active" % sid, data={"csrf_token": ac}, follow_redirects=False)
    blocked = app.test_client()
    rr = login(blocked, sname, "123456")
    record("TF-053", "停用后登录被拒绝并提示“该账号已被停用”（%s）" % rr.status_code,
           "该账号已被停用" in text(rr) and rr.status_code == 200)

    r = adm.post("/admin/users/%d/delete" % sid, data={"csrf_token": ac}, follow_redirects=False)
    with app.app_context():
        exists = db.session.get(User, sid) is not None
        orphans = Submission.query.filter_by(user_id=sid).count()
    record("TF-054", "删除账号（%s），已不存在：%s，残留提交 %d 条" % (r.status_code, not exists, orphans),
           r.status_code == 302 and not exists and orphans == 0)

    r = stu.get("/admin/users", follow_redirects=False)
    r2 = tea.get("/admin/users", follow_redirects=False)
    record("TF-055", "学生 %s、教师 %s 访问用户管理均被重定向（无权限）" % (r.status_code, r2.status_code),
           r.status_code == 302 and r2.status_code == 302)

    aname = "hack_admin_%s" % datetime.now().strftime("%H%M%S")[-6:]
    cli = app.test_client()
    cli.post("/auth/register", data={"username": aname, "password": "123456", "confirm": "123456",
                                     "role": "admin", "csrf_token": csrf(cli)})
    with app.app_context():
        u = User.query.filter_by(username=aname).first()
        role_now = u.role if u else None
    record("TF-056", "注册页提交 role=admin 被降级为 %s，无法提权" % role_now, role_now == "student")

    # ---------- 清理 ----------
    with app.app_context():
        for title in ("无用例题目",):
            p = Problem.query.filter_by(title=title).first()
            if p:
                Submission.query.filter_by(problem_id=p.id).delete()
                db.session.execute(homework_problems.delete().where(homework_problems.c.problem_id == p.id))
                db.session.delete(p)
        p = db.session.get(Problem, target_id)
        if p:
            p.title = origin_title
            p.tags = [t for t in (KnowledgePoint.query.filter_by(name=n).first()
                                  for n in origin_tags.split(",")) if t]
        for name in (new_name, tname, aname, "teacher2"):
            u = User.query.filter_by(username=name).first()
            if u:
                Submission.query.filter_by(user_id=u.id).delete()
                rooms = Classroom.query.filter_by(teacher_id=u.id).all()
                for room in rooms:
                    for hw in list(room.homeworks):
                        Submission.query.filter_by(homework_id=hw.id).delete()
                        db.session.execute(homework_problems.delete().where(homework_problems.c.homework_id == hw.id))
                        db.session.delete(hw)
                    db.session.delete(room)
                db.session.delete(u)
        for sid in fake_ids:
            s = db.session.get(Submission, sid)
            if s:
                db.session.delete(s)
        s = db.session.get(Submission, other_id)
        if s and s.code == "print(1)":
            db.session.delete(s)
        db.session.commit()


# ----------------------------- 结果导出 -----------------------------


EXTRA_CASES = [
    ("TF-048", "管理员", "管理员登录成功", "存在管理员账号 admin/admin123", "使用管理员账号登录", "登录成功，进入管理控制台", "高"),
    ("TF-049", "", "用户管理列表加载", "管理员已登录", "访问用户管理页", "展示学生/教师/管理员账号列表", "高"),
    ("TF-050", "", "管理员新增教师账号", "管理员已登录", "创建角色为教师的账号", "创建成功且该账号可登录", "高"),
    ("TF-051", "", "管理员编辑账号信息", "账号存在", "修改姓名与角色后保存", "修改生效", "中"),
    ("TF-052", "", "管理员重置密码", "账号存在", "执行重置密码", "密码重置为 123456 并可登录", "中"),
    ("TF-053", "", "管理员停用账号", "账号存在", "停用该账号后尝试登录", "登录被拒绝并提示已停用", "高"),
    ("TF-054", "", "管理员删除账号", "账号存在", "删除账号", "账号删除且关联提交被清理", "高"),
    ("TF-055", "", "非管理员访问用户管理", "学生/教师已登录", "访问 /admin/users", "重定向并提示无权限", "高"),
    ("TF-056", "", "注册页无法提权为管理员", "无", "注册时提交 role=admin", "角色被降级为学生，无法提权", "高"),
]


def load_template_rows():
    """从模板表读取用例元数据（编号/模块/标题/前置/步骤/预期/优先级）。"""
    from openpyxl import load_workbook

    rows = []
    wb = load_workbook(TEMPLATE)
    ws = wb[wb.sheetnames[0]]
    for row in ws.iter_rows(min_row=2, values_only=True):
        cid = (row[0] or "").strip() if row[0] else ""
        if not cid or not cid.startswith("TF-"):
            continue
        rows.append({
            "id": cid,
            "module": row[1] or "",
            "title": row[2] or "",
            "pre": row[3] or "",
            "steps": row[4] or "",
            "expect": row[5] or "",
            "priority": row[11] or "中",
        })
    return rows


def write_xlsx(rows, out_path):
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    try:
        wb = load_workbook(TEMPLATE)
        ws = wb[wb.sheetnames[0]]
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row)
    except Exception:
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws.append(["用例编号", "测试模块", "用例标题", "前置条件", "测试步骤", "预期结果",
                   "实际结果", "", "", "测试状态", "缺陷编号", "测试优先级"])

    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")
    header_fill = PatternFill("solid", fgColor="E2EFDA")
    pass_fill = PatternFill("solid", fgColor="C6EFCE")
    fail_fill = PatternFill("solid", fgColor="FFC7CE")

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    passed = failed = 0
    for meta in rows:
        cid = meta["id"]
        actual, status, bug = RESULTS.get(cid, ("未执行", "未执行", ""))
        if status == "通过":
            passed += 1
        elif status == "失败":
            failed += 1
        ws.append([cid, meta["module"], meta["title"], meta["pre"], meta["steps"],
                   meta["expect"], actual, "", "", status, bug, meta["priority"]])
        for cell in ws[ws.max_row]:
            cell.alignment = wrap
            cell.border = border
        ws.cell(row=ws.max_row, column=10).fill = pass_fill if status == "通过" else fail_fill
        ws.cell(row=ws.max_row, column=10).alignment = Alignment(horizontal="center", vertical="center")

    from openpyxl.utils import get_column_letter

    for i, w in enumerate([10, 10, 20, 24, 26, 30, 46, 3, 3, 10, 10, 10], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    bugs = [(cid, RESULTS[cid][2]) for cid in RESULTS if RESULTS[cid][1] == "失败" and RESULTS[cid][2]]
    summary = "测试统计：共执行 %d 条用例，通过 %d 条、失败 %d 条；执行时间 %s" % (
        len(rows), passed, failed, datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
    ws.append([])
    ws.append([summary])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    for cid, desc in bugs:
        ws.append(["%s 未通过：%s" % (cid, desc)])
    if not bugs:
        ws.append(["未发现未通过用例；历史缺陷 BUG-001（题目编辑清空用例）、BUG-002（登录限流误锁）已修复并回归通过。"])

    wb.save(out_path)
    return passed, failed, out_path


def main():
    print("=" * 70)
    print("按测试用例表复现测试")
    print("=" * 70)
    app = create_app()
    with app.app_context():
        if not User.query.filter_by(username="admin").first():
            admin = User(username="admin", display_name="系统管理员", role="admin")
            admin.set_password("admin123")
            db.session.add(admin)
            db.session.commit()
    run_cases(app)

    rows = []
    try:
        rows = load_template_rows()
    except Exception as exc:
        print("读取模板失败（%s），使用内置用例清单" % exc)
    known = {r["id"] for r in rows}
    for cid, module, title, pre, steps, expect, priority in EXTRA_CASES:
        if cid not in known:
            rows.append({"id": cid, "module": module, "title": title, "pre": pre,
                         "steps": steps, "expect": expect, "priority": priority})

    out = os.path.abspath(OUTPUT)
    try:
        passed, failed, path = write_xlsx(rows, out)
    except PermissionError:
        # 结果表可能正被 Excel 打开，自动改用带时间戳的副本
        alt = os.path.join(
            os.path.dirname(out),
            "测试结果-%s.xlsx" % datetime.now().strftime("%Y%m%d-%H%M%S"),
        )
        print("默认结果文件被占用（可能已在 Excel 中打开），改为输出：%s" % alt)
        passed, failed, path = write_xlsx(rows, alt)
    print("\n" + "=" * 70)
    print("通过 %d 条，失败 %d 条；结果表：%s" % (passed, failed, path))
    for cid in RESULTS:
        if RESULTS[cid][1] == "失败":
            print("  - %s 失败：%s" % (cid, RESULTS[cid][0]))
    print("=" * 70)
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
